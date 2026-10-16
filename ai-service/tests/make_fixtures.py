"""
生成解析器测试样本（PDF / DOCX）
================================
Markdown 之外，文档中心还声明支持 PDF 与 Word。
这两个解析路径必须有可复现的测试样本，否则「支持 PDF/Word」就是空话。

运行：
    cd ai-service
    pip install fpdf2            # 仅用于生成测试 PDF，不进 requirements
    python tests/make_fixtures.py

产物（data/samples/）：
    宿舍与后勤服务.docx      —— 由 python-docx 生成，含段落 + 表格
    图书馆使用手册.pdf       —— 由 fpdf2 生成，多页，用于验证**页码溯源**
"""
from __future__ import annotations

import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLES = os.path.join(BASE_DIR, "data", "samples")
os.makedirs(SAMPLES, exist_ok=True)

# Windows 中文字体，fpdf2 需要显式注册 TTF
CN_FONT = r"C:\Windows\Fonts\msyh.ttc"


def make_docx() -> str:
    """生成 Word 样本：含标题、段落与一张表格。"""
    import docx
    from docx.shared import Pt

    document = docx.Document()
    document.add_heading("宿舍与后勤服务手册", level=1)

    document.add_heading("一、宿舍用电", level=2)
    document.add_paragraph(
        "学生宿舍实行定额用电制度，每位学生每月享有 8 度免费额度，"
        "超出部分按照 0.55 元/度计费，可在「一卡通」小程序中自助充值。"
    )
    document.add_paragraph(
        "严禁使用大功率电器（电热毯、电磁炉、电饭锅等），"
        "一经发现将没收电器并给予通报批评，情节严重者取消住宿资格。"
    )

    document.add_heading("二、宿舍报修", level=2)
    document.add_paragraph(
        "宿舍设施损坏可通过「烟理工后勤」公众号提交报修单，"
        "填写宿舍号与故障描述并上传照片，一般 24 小时内上门处理。"
    )
    document.add_paragraph("紧急情况（漏水、断电、门锁损坏）请直接拨打后勤值班电话 0535-6915xxx。")

    document.add_heading("三、后勤服务一览", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "服务项目"
    hdr[1].text = "办理地点"
    hdr[2].text = "服务时间"
    for row in [
        ("快递代收", "二食堂东侧菜鸟驿站", "8:30 - 20:30"),
        ("洗衣房", "每栋宿舍楼一层", "6:00 - 23:00"),
        ("开水房", "宿舍楼每层两端", "全天开放"),
        ("电费充值", "一卡通小程序", "24 小时在线"),
    ]:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v

    path = os.path.join(SAMPLES, "宿舍与后勤服务.docx")
    document.save(path)
    print(f"  已生成 DOCX: {path}  ({os.path.getsize(path) / 1024:.1f} KB)")
    return path


def make_pdf() -> str:
    """
    生成 PDF 样本：**故意做成多页**，用来验证解析后页码是否正确保留。
    fpdf2 走的是标准 PDF 文本对象，pypdf 能正常抽取（扫描版才需要 OCR）。
    """
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)

    if not os.path.exists(CN_FONT):
        raise FileNotFoundError(f"缺少中文字体：{CN_FONT}")
    pdf.add_font("msyh", "", CN_FONT)      # 注册中文字体，否则中文全是方块/乱码

    pages = [
        ("图书馆使用手册", [
            "本手册介绍图书馆的开放时间、借阅规则、电子资源访问方式与自习座位管理办法。",
            "图书馆是学校文献信息中心，入馆请刷校园卡或出示学生证，保持安静，禁止携带食物入内。",
            "一、开放时间",
            "周一至周五：8:00 - 22:00；周六、周日：9:00 - 21:00。",
            "寒暑假及法定节假日另行通知，一般调整为 9:00 - 17:00。",
            "闭馆前 15 分钟开始清场，请读者提前办理借还手续。",
            "每周三下午 14:00 - 17:00 为馆内整理时间，部分书库暂停开放。",
            "如遇考试周，图书馆将延长开放时间至 23:00，具体安排见馆内公告。",
        ]),
        ("二、借阅规则", [
            "本科生每次最多可借 10 册，借期 30 天；研究生每次最多可借 20 册，借期 60 天。",
            "可续借 2 次，每次 15 天，续借需在到期前通过图书馆微信公众号或自助借还机办理。",
            "逾期未还将暂停借阅权限，并按每天 0.2 元收取滞纳金，单册上限 20 元。",
            "所借图书遗失的，须购买相同版本图书赔偿，并缴纳 5 元加工费。",
            "工具书、期刊、珍本图书仅限馆内阅览，一律不外借。",
            "毕业生离校前须归还全部借阅图书，否则不予办理离校手续。",
            "读者可通过馆藏检索系统查询图书位置，索书号由中图法分类号加著者号组成。",
        ]),
        ("三、电子资源", [
            "校园网内可直接访问中国知网、万方、维普、IEEE、Springer 等中外文数据库，无需登录。",
            "校外访问需先登录 WebVPN，地址为 vpn.yitsd.edu.cn。",
            "WebVPN 登录账号为学号，初始密码为统一身份认证密码，首次登录后请及时修改。",
            "如遇账号异常或无法访问，请联系图书馆技术部 0535-6915xxx。",
            "文献传递服务：本馆未收藏的文献可提交文献传递申请，一般 2 个工作日内发送至邮箱。",
            "校外访问单次会话有效期为 8 小时，超时需重新登录。",
        ]),
        ("四、自习座位", [
            "图书馆三至五层为自习区，需通过座位预约系统提前预约，可预约当日及次日座位。",
            "预约后需在 30 分钟内刷校园卡入座，否则座位自动释放并记一次违约。",
            "爽约 3 次将冻结预约权限 7 天，累计 10 次冻结 30 天。",
            "自习区禁止占座，离开超过 60 分钟且未办理暂离的，工作人员有权清理桌面物品。",
            "馆内提供 12 间研讨间，需 3 人以上方可预约，单次使用不超过 4 小时。",
        ]),
    ]

    for title, lines in pages:
        pdf.add_page()
        pdf.set_font("msyh", size=18)
        pdf.cell(0, 12, title, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
        pdf.set_font("msyh", size=12)
        for line in lines:
            pdf.multi_cell(0, 8, line)
            pdf.ln(1)

    path = os.path.join(SAMPLES, "图书馆使用手册.pdf")
    pdf.output(path)
    print(f"  已生成 PDF : {path}  ({os.path.getsize(path) / 1024:.1f} KB, {len(pages)} 页)")
    return path


if __name__ == "__main__":
    print("生成测试样本 ...")
    make_docx()
    make_pdf()
    print("完成。")
